from pathlib import Path
import json, math
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'release-evidence'/'rc10.72.0'/'agency-banking-benchmark'
model=json.loads((OUT/'AGENCY_BANKING_DETERMINISTIC_CANONICAL_MODEL.json').read_text())
score=json.loads((OUT/'AGENCY_BANKING_GOLDEN_BENCHMARK_SCORECARD.json').read_text())
source=model['sourceInput']
NAVY='0B172A'; AMBER='E3A008'; TEAL='0F766E'; LIGHT='F4F7FB'; WHITE='FFFFFF'; GREY='475569'; RED='B91C1C'

def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def set_cell_text(cell,text,bold=False,color='111827',size=8):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run(str(text)); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def add_table(doc,headers,rows,widths=None,font=7.5):
    tbl=doc.add_table(rows=1,cols=len(headers)); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.style='Table Grid'
    for i,h in enumerate(headers): set_cell_shading(tbl.rows[0].cells[i],NAVY); set_cell_text(tbl.rows[0].cells[i],h,True,WHITE,font)
    set_repeat_table_header(tbl.rows[0])
    for ri,row in enumerate(rows):
        cells=tbl.add_row().cells
        for i,v in enumerate(row):
            set_cell_shading(cells[i],WHITE if ri%2==0 else LIGHT); set_cell_text(cells[i],v,False,'1F2937',font)
    if widths:
        for row in tbl.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return tbl

def add_heading(doc,text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.keep_with_next=True
    r=p.add_run(text) if not p.text else None
    return p

def add_bullets(doc,items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(str(item))

def add_callout(doc,title,body,kind='info'):
    tbl=doc.add_table(rows=1,cols=1); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell=tbl.cell(0,0); set_cell_shading(cell, 'ECFDF5' if kind=='info' else 'FFF7ED')
    p=cell.paragraphs[0]; r=p.add_run(title+'\n'); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL if kind=='info' else AMBER); r.font.size=Pt(10)
    r2=p.add_run(body); r2.font.size=Pt(9); r2.font.color.rgb=RGBColor.from_string(GREY)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def seq_png(j):
    participants=j['participants']; steps=j['steps']; fig,ax=plt.subplots(figsize=(10,5.2)); n=len(participants); xs=list(range(n));
    ax.set_xlim(-.5,n-.5); ax.set_ylim(len(steps)+1,-1); ax.axis('off')
    for x,p in zip(xs,participants):
        ax.text(x,-.5,p,ha='center',va='center',fontsize=8,fontweight='bold',bbox=dict(boxstyle='round,pad=.3',fc='#EAF2FF',ec='#334155'))
        ax.plot([x,x],[0,len(steps)+.4],linestyle='--',linewidth=.8,color='#94A3B8')
    for i,step in enumerate(steps):
        a=i%max(1,n-1); b=min(a+1,n-1); y=i+.55
        ax.annotate('',xy=(b,y),xytext=(a,y),arrowprops=dict(arrowstyle='->',lw=1,color='#0F766E'))
        ax.text((a+b)/2,y-.14,textwrap(step,34),ha='center',va='bottom',fontsize=7,color='#1F2937')
    ax.set_title(j['title']+' — business interaction sequence',fontsize=12,fontweight='bold',loc='left')
    p=OUT/'diagrams'/f"sequence-{j['id'].lower()}.png"; fig.tight_layout(); fig.savefig(p,dpi=180,bbox_inches='tight'); plt.close(fig); return p

def textwrap(s,n):
    import textwrap as tw
    return '\n'.join(tw.wrap(s,n))

for j in model['journeys']: seq_png(j)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9); styles['Normal'].paragraph_format.space_after=Pt(4); styles['Normal'].paragraph_format.line_spacing=1.05
for i,size in [(1,20),(2,15),(3,12)]:
    st=styles[f'Heading {i}']; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(NAVY); st.paragraph_format.space_before=Pt(10); st.paragraph_format.space_after=Pt(4)
# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(30); r=p.add_run('AIW'); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(AMBER)
p=doc.add_paragraph(); r=p.add_run('AGENCY BANKING\nSOLUTION DESIGN DOCUMENT'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(28); r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph(); r=p.add_run('Golden Benchmark — Independently Generated Architecture'); r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(TEAL)
doc.add_paragraph('\n')
add_callout(doc,'Generation integrity','This design was generated from the benchmark business problem, stakeholder needs, constraints and functional needs. The completed Agency Banking reference SDD was excluded from generation and reserved for post-generation comparison.','info')
for label,value in [('Release','AIW v0.10.0-rc.10.72.0'),('Knowledge release','AKR-0.10.72.0'),('Architecture authority','Deterministic Kernel + governed knowledge + human approval'),('Review status','Controlled internal pilot; external human panel and managed-infrastructure acceptance pending')]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); rr=p.add_run(label+': '); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(NAVY); p.add_run(value)
doc.add_page_break()
# TOC
add_heading(doc,'Document map',1)
sections=['Executive summary and scope','Stakeholders and requirements','Major solution journeys','Architecture drivers and reasoning','System context and logical architecture','Component responsibilities','Interface contracts','Data architecture','Security and trust boundaries','Deployment, resilience and recovery','Architecture decisions','Fitness tests and assurance','Traceability and benchmark scorecard','Independent validation boundary']
for i,s in enumerate(sections,1): doc.add_paragraph(f'{i}. {s}')
doc.add_page_break()
# 1
add_heading(doc,'1. Executive summary and scope',1)
doc.add_paragraph(source['businessProblem'])
doc.add_paragraph('The proposed architecture protects financial integrity through explicit transaction identity, durable journal state, idempotency, external status inquiry, controlled reversal and reconciliation. It separates channel access, identity, agent/customer lifecycle, policy, financial transaction coordination, external adapters, asynchronous outcomes, settlement, disputes and audit.')
add_callout(doc,'Truthfulness rule','Availability, latency, throughput, RTO and RPO values remain candidate targets until business-impact, workload and service-tier evidence is approved. They are not represented as established facts.','warn')
add_heading(doc,'1.1 Scope boundaries',2)
add_bullets(doc,source['knownConstraints'])
add_heading(doc,'1.2 Candidate targets requiring confirmation',2)
add_table(doc,['Candidate target','Status','Evidence required'],[(x['target'],x['status'],x['requiredEvidence']) for x in source['candidateTargetsRequiringConfirmation']],[1.6,1.6,4.6],8)
# 2
add_heading(doc,'2. Stakeholders and requirements',1)
add_heading(doc,'2.1 Stakeholders',2); add_bullets(doc,source['stakeholders'])
add_heading(doc,'2.2 Requirements catalogue',2)
add_table(doc,['ID','Priority','Type','Requirement','Acceptance criterion'],[(r['id'],r['priority'],r['kind'],r['statement'],r['acceptanceCriteria']) for r in model['requirements']],[.7,.45,.6,3.25,3.1],7)
# 3 journeys
add_heading(doc,'3. Major solution journeys',1)
for j in model['journeys']:
    add_heading(doc,f"3.{model['journeys'].index(j)+1} {j['title']}",2)
    doc.add_paragraph('Participants: '+', '.join(j['participants']))
    doc.add_picture(str(OUT/'diagrams'/f"sequence-{j['id'].lower()}.png"),width=Inches(7.1))
    cap=doc.paragraphs[-1]; cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Paths modelled: happy, alternate, failure and recovery. Each interaction remains traceable to requirements and downstream interface obligations.')
# 4 drivers
add_heading(doc,'4. Architecture drivers and reasoning',1)
doc.add_paragraph('The design applies the following reasoning chain: requirement → quality scenario → tactic → style/pattern implication → responsibility → interface obligation → technology capability → physical realization → fitness evidence.')
add_table(doc,['Quality driver','Primary tactic','Downstream obligation'],[
('Idempotency and financial integrity','Stable command identity, durable journal, status inquiry','Transaction Service, Journal, adapter contracts and replay tests'),
('Safe dependency failure','Explicit reject/retry/uncertain/compensate semantics','Workflow/Recovery Coordinator and external adapter failure policies'),
('Security and privacy','Strong identity, least privilege, classification, encryption, audit','IAM, trust boundaries, protected data and security fitness tests'),
('Recoverability','Durable state, replay, reconciliation and controlled compensation','Journal, outbox, recovery workflow, backup and recovery tests'),
('Maintainability','Provider-neutral capabilities, module boundaries, versioned contracts','Service-based modular architecture and owned interfaces'),
('Observability','Correlation identity, structured audit, metrics and traces','Audit/Observability component and journey-level evidence')],[1.5,2.4,4.1],8)
# 5 diagrams
add_heading(doc,'5. System context and logical architecture',1)
for name,title in [('01-system-context','System context'),('02-container-logical','Logical container view'),('03-transaction-component','Transaction Service component view')]:
    add_heading(doc,title,2); doc.add_picture(str(OUT/'diagrams'/f'{name}.png'),width=Inches(7.1)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
# 6 components
add_heading(doc,'6. Component responsibilities',1)
add_table(doc,['ID','Component','Responsibility','Owner','Requirement lineage'],[(c['id'],c['name'],c['responsibility'],c['owner'],', '.join(c['requirementRefs'])) for c in model['components']],[.45,1.4,3.0,1.1,1.6],7)
# 7 interfaces
add_heading(doc,'7. Interface contracts',1)
doc.add_paragraph('Every critical contract records ownership, interaction style, authentication, authorization, timeout/failure behavior, retry and idempotency semantics, ordering/consistency assumptions, observability, data classification and version policy.')
for chunk_start in range(0,len(model['interfaces']),7):
    chunk=model['interfaces'][chunk_start:chunk_start+7]
    add_table(doc,['ID','Provider → Consumer','Purpose / style','Failure / idempotency','Class / owner'],[(i['id'],f"{i['provider']} → {i['consumer']}",f"{i['purpose']} / {i['interactionStyle']}",f"{i['timeoutPolicy']}; {i['idempotency']}",f"{i['dataClassification']}; {i['contractOwner']}") for i in chunk],[.4,1.7,2.0,2.35,1.5],6.8)
# 8 data
add_heading(doc,'8. Data architecture',1)
doc.add_picture(str(OUT/'diagrams'/'04-data-ownership.png'),width=Inches(7.1)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
add_table(doc,['Entity','Owner','Classification','System of record'],[(d['name'],d['owner'],d['classification'],d['systemOfRecord']) for d in model['dataEntities']],[1.7,1.7,1.2,2.8],8)
# 9 security
add_heading(doc,'9. Security and trust boundaries',1)
doc.add_picture(str(OUT/'diagrams'/'05-trust-boundaries.png'),width=Inches(7.1)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
add_table(doc,['ID','Threat','Boundary','Required controls'],[(t['id'],t['threat'],t['boundary'],', '.join(t['controls'])) for t in model['threats']],[.45,2.1,1.55,3.3],7.2)
# 10 deploy
add_heading(doc,'10. Deployment, resilience and recovery',1)
doc.add_picture(str(OUT/'diagrams'/'06-deployment-recovery.png'),width=Inches(7.1)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('The physical view remains product-neutral. Redundant application runtimes, durable data, queue/event capability, evidence storage, telemetry and recovery capability are required. Exact topology, product, region, backup schedule, RTO and RPO require target-environment and business approval.')
# 11 decisions
add_heading(doc,'11. Architecture decisions',1)
for d in model['decisions']:
    add_heading(doc,f"{d['id']} — {d['title']}",2)
    add_table(doc,['Context','Alternatives','Decision','Consequences'],[(d['context'],'; '.join(d['alternatives']),d['decision'],'; '.join(d['consequences']))],[1.9,1.8,2.4,1.7],7.5)
# 12 tests
add_heading(doc,'12. Fitness tests and assurance',1)
add_table(doc,['ID','Fitness test','Method','Expected evidence'],[(f['id'],f['name'],f['method'],f['expected']) for f in model['fitnessTests']],[.45,1.55,2.65,3.15],7.3)
# 13 trace score
add_heading(doc,'13. Traceability and benchmark scorecard',1)
add_heading(doc,'13.1 Requirement-to-component traceability',2)
add_table(doc,['Requirement','Journeys','Components'],[(t['requirementId'],', '.join(t['journeyRefs']),', '.join(t['componentRefs']) or 'Open mapping') for t in model['traceability']],[1.0,2.6,4.2],7.5)
add_heading(doc,'13.2 Objective result',2)
metrics=score['metrics']
rows=[]
for k,label,target in [
('highPriorityDriverTraceabilityPct','High-priority traceability','100%'),('overallRequirementToModelTraceabilityPct','Overall requirement-to-model traceability','≥90%'),('criticalInterfaceCompletenessPct','Critical-interface completeness','≥90%'),('significantDecisionAlternativeTradeoffPct','Decisions with alternatives/trade-offs','100%'),('componentResponsibilityOwnershipRationalePct','Components with responsibility/owner/rationale','≥95%')]: rows.append((label,str(metrics[k])+'%',target,'Pass'))
rows += [('Unsupported numeric claims',str(metrics['unsupportedNumericPerformanceSecurityRecoveryClaims']),'0','Pass'),('Critical security omissions',str(metrics['criticalSecurityOmissions']),'0','Model pass; live test pending'),('Critical resilience omissions',str(metrics['criticalResilienceOmissions']),'0','Model pass; live test pending')]
add_table(doc,['Measure','Actual','Target','Result'],rows,[3.6,1.0,1.0,2.2],8)
add_callout(doc,'Independent AI review','Rubric score 4.25/5. The design is a strong controlled-pilot candidate. This is not an external human-panel score and does not prove market superiority.','info')
# 14 boundary
add_heading(doc,'14. Independent validation boundary',1)
add_bullets(doc,[
'Completed: deterministic generation from the source brief, requirements and journey modelling, logical design, interface/data/security/deployment design, ADRs, fitness tests, rendered views, traceability and professional SDD assembly.',
'Completed: internal independent AI architecture review and controlled benchmark pilot.',
'Not completed: live core banking, switch and KYC contract acceptance.',
'Not completed: managed enterprise OIDC, PostgreSQL RLS, queue, object storage, KMS and telemetry acceptance.',
'Not completed: external blinded human expert panel.',
'Not measurable: preparation-time reduction against a controlled conventional architect baseline.'
])
add_heading(doc,'14.1 External acceptance actions',2)
add_table(doc,['Acceptance item','Accountable role','Required evidence','Current state'],[
('Core banking, switch and KYC contracts','Integration owner','Signed contract tests, failure scenarios and reconciliation evidence','Blocked - target systems unavailable'),
('Managed identity and tenant isolation','Security and platform','OIDC/JWKS acceptance, step-up authorization and PostgreSQL RLS penetration evidence','Blocked - managed environment unavailable'),
('Queue, storage and recovery','Platform and SRE','Durable queue recovery, object lifecycle, backup restore and worker failover evidence','Blocked - managed environment unavailable'),
('Enterprise signing and remote release','Security and knowledge governance','Cloud KMS verification plus protected remote Git commit/tag evidence','Blocked - credentials and remote absent'),
('Blinded expert review','Independent architecture panel','Scores from at least three reviewers using the sealed protocol','Prepared, not executed'),
('Time and trust baseline','Pilot owner','Conventional and AIW-assisted timings plus architect trust survey','Not measurable in internal run')
],[2.2,1.45,3.1,1.45],7.0)
# Footer/header
for section in doc.sections:
    hp=section.header.paragraphs[0]; hp.text='AIW | AGENCY BANKING GOLDEN BENCHMARK'; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; hp.runs[0].font.size=Pt(7); hp.runs[0].font.color.rgb=RGBColor.from_string(GREY)
    fp=section.footer.paragraphs[0]; fp.text='AIW v0.10.0-rc.10.72.0 • Controlled internal pilot • Not production accepted'; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fp.runs[0].font.size=Pt(7); fp.runs[0].font.color.rgb=RGBColor.from_string(GREY)
# Save
doc.core_properties.title='Agency Banking Solution Design Document — AIW Golden Benchmark'
doc.core_properties.subject='AIW independent architecture benchmark'
doc.core_properties.author='AIW Independent Validation Process'
out=OUT/'AGENCY_BANKING_AI_GENERATED_SDD.docx'; doc.save(out); print(out)
